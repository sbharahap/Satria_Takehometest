"""Paired local-only model comparison. No cloud calls or production data mutations.

Run: uv run python -m scripts.compare_local_models --repeats 2
Generator tests disable semantic verification to avoid one candidate judging another.
Verifier tests independently use the same fixed true/false claims for each model.
Automatic regex/source checks are transparent heuristics, not semantic accuracy.
"""

import argparse
import hashlib
import io
import json
import logging
import re
import statistics
import time
from uuid import uuid4

import httpx
from docx import Document
from pypdf import PdfReader

from src.bootstrap import initialize
from src.config import PROJECT_ROOT, Settings
from src.knowledge.embeddings import LocalEmbeddings
from src.models import AgentInput, Citation
from src.providers.base import ChatProvider, ChatRequest
from src.providers.litellm_adapter import LiteLLMProvider
from src.providers.registry import ProviderRegistry
from src.services.chat import UNKNOWN, UNVERIFIED
from src.services.verification import CitationVerifier

MODELS = ["ollama_chat/qwen3:8b", "ollama_chat/llama3.1:8b"]


def assess(result, case):
    answer = result["content"]
    citations = result["citations"]
    if case.get("unknown"):
        return {
            "pass": answer == UNKNOWN and not citations,
            "strict_unknown": answer == UNKNOWN,
            "no_citations": not citations,
        }
    text = answer.casefold()
    expected = any(case["file"] in c["filename"] and c["page"] in case["pages"] for c in citations)
    keywords = all(re.search(pattern, text) for pattern in case.get("required", []))
    forbidden = any(re.search(pattern, text) for pattern in case.get("forbidden", []))
    return {
        "pass": bool(citations and expected and keywords and not forbidden),
        "answered": bool(citations),
        "expected_citation": expected,
        "keyword_rubric": bool(keywords and not forbidden),
    }


class RecordingProvider(ChatProvider):
    def __init__(self, settings):
        self.provider = LiteLLMProvider(settings)
        self.calls = []

    def complete(self, request):
        start = time.perf_counter()
        try:
            response = self.provider.complete(request)
            self.calls.append(
                {
                    "seconds": time.perf_counter() - start,
                    "response": response.content,
                    "phase": "rewrite" if request.max_tokens == 300 else "generate",
                }
            )
            return response
        except Exception as error:
            self.calls.append(
                {"seconds": time.perf_counter() - start, "error_type": type(error).__name__}
            )
            raise


class TraceGraph:
    def __init__(self, graph):
        self.graph, self.last = graph, {}

    def invoke(self, state):
        self.last = self.graph.invoke(state)
        return self.last


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repeats", type=int, default=2, choices=range(1, 4))
    parser.add_argument(
        "--answer-mode", choices=["extractive", "abstractive"], default="extractive"
    )
    args = parser.parse_args()
    settings = Settings()
    cases = json.loads((PROJECT_ROOT / "eval/local_model_cases.json").read_text())["cases"]
    root = settings.data_dir / "model-comparison" / uuid4().hex
    root.mkdir(parents=True)
    documents = {
        group: sorted((settings.data_dir / "demo" / group).glob("*_id.pdf"))
        for group in ("hr", "legal")
    }
    assert all(len(files) == 2 for files in documents.values()), "Run scripts.prepare_demo first."
    report = {
        "protocol": {
            "models": MODELS,
            "answer_mode": args.answer_mode,
            "source_sha256": {
                str(path): hashlib.sha256((PROJECT_ROOT / path).read_bytes()).hexdigest()
                for path in ["src/services/chat.py", "src/knowledge/grounding.py", "src/config.py"]
            },
            "repeats": args.repeats,
            "temperature": 0,
            "context_tokens": 8192,
            "max_output_tokens": 200 if args.answer_mode == "extractive" else 700,
            "semantic_verifier_for_generation": False,
            "cloud_requests": 0,
            "case_file_sha256": hashlib.sha256(
                (PROJECT_ROOT / "eval/local_model_cases.json").read_bytes()
            ).hexdigest(),
            "model_order": "alternating per repeat",
            "warmup_excluded_from_case_latency": True,
            "limitations": "Small policy-specific heuristic rubric, no general robustness ranking.",
        },
        "document_sha256": {
            p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for files in documents.values()
            for p in files
        },
        "model_details": {},
        "warmups": [],
        "records": [],
        "verifier_records": [],
    }
    with httpx.Client(timeout=20) as client:
        tags = (
            client.get(settings.ollama_base_url + "/api/tags").raise_for_status().json()["models"]
        )
        for model in MODELS:
            name = model.split("/", 1)[1]
            entry = next(item for item in tags if item["name"] == name)
            report["model_details"][model] = {
                k: entry.get(k) for k in ("digest", "size", "details")
            }
    embeddings = LocalEmbeddings(settings)

    def save():
        (root / "results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")

    print(f"REPORT: {root / 'results.json'}", flush=True)
    for repeat in range(args.repeats):
        for model in MODELS if repeat % 2 == 0 else list(reversed(MODELS)):
            run_settings = settings.model_copy(
                update={
                    "data_dir": root / f"r{repeat + 1}-{model.split('/')[-1].replace(':', '-')}",
                    "citation_verifier_enabled": False,
                    "answer_mode": args.answer_mode,
                    "local_fallback_model": model,
                }
            )
            provider = RecordingProvider(run_settings)
            start = time.perf_counter()
            try:
                provider.complete(
                    ChatRequest(
                        model,
                        [{"role": "user", "content": 'Return JSON {"ok":true}.'}],
                        max_tokens=32,
                        json_mode=True,
                    )
                )
                warmup_error = None
            except Exception as error:
                warmup_error = type(error).__name__
            report["warmups"].append(
                {
                    "model": model,
                    "repeat": repeat + 1,
                    "seconds": time.perf_counter() - start,
                    "error": warmup_error,
                }
            )
            services = initialize(run_settings)
            services.documents.embeddings = services.chat.embeddings = embeddings
            services.chat.providers.register(
                "litellm", lambda settings, provider=provider: provider
            )
            services.chat.graph = TraceGraph(services.chat.graph)
            agents = {}
            for group, paths in documents.items():
                agents[group] = services.agents.create(
                    AgentInput(
                        name=group,
                        instructions="Jawab berdasarkan dokumen aktif. Jangan mengarang kebijakan.",
                        provider="litellm",
                        model=model,
                    )
                )
                for path in paths:
                    doc = services.documents.upload(agents[group].id, path.name, path.read_bytes())
                    assert doc.status == "ready", doc.error_message

            def ask(
                case,
                agent=None,
                conversation=None,
                *,
                agents=agents,
                services=services,
                provider=provider,
                model=model,
                repeat=repeat,
            ):
                agent = agent or agents[case["agent"]]
                conversation = conversation or services.chat.new_conversation(agent.id)
                call_start = len(provider.calls)
                start = time.perf_counter()
                record = {
                    "model": model,
                    "repeat": repeat + 1,
                    "id": case["id"],
                    "agent": case["agent"],
                    "question": case["question"],
                    "expected_unknown": case.get("unknown", False),
                }
                try:
                    answer = services.chat.ask(agent.id, conversation, case["question"])
                    record.update(
                        answer=answer["content"],
                        citations=answer["citations"],
                        checks=assess(answer, case),
                    )
                    state = services.chat.graph.last
                    sources = [
                        {"filename": s.filename, "page": s.page, "text": s.text}
                        for s in state.get("sources", [])
                    ]
                    record["query"] = state.get("query")
                    record["retrieved_sources"] = sources
                    record["retrieval_fingerprint"] = hashlib.sha256(
                        json.dumps(sources, sort_keys=True, ensure_ascii=False).encode()
                    ).hexdigest()
                    record["retrieval_gold_hit"] = (
                        None
                        if case.get("unknown")
                        else any(
                            case["file"] in s["filename"] and s["page"] in case["pages"]
                            for s in sources
                        )
                    )
                except Exception as error:
                    record.update(error_type=type(error).__name__, checks={"pass": False})
                record["seconds"] = time.perf_counter() - start
                record["provider_calls"] = provider.calls[call_start:]
                report["records"].append(record)
                save()
                print(
                    json.dumps(
                        {k: record[k] for k in ("model", "repeat", "id", "checks", "seconds")}
                    ),
                    flush=True,
                )
                return record

            for case in cases:
                ask(case)

            # Native DOCX is a format-conversion fixture from public policy, not a corporate DOCX.
            text = (
                PdfReader(next(p for p in documents["hr"] if "people_policy" in p.name))
                .pages[2]
                .extract_text()
            )
            excerpt = text.split("6. Tata Kelola Karyawan", 1)[1].split("7. Ketentuan Umum", 1)[0]
            word, buffer = Document(), io.BytesIO()
            word.add_heading("Tata Kelola Karyawan", 1)
            word.add_paragraph(" ".join(excerpt.split()))
            word.save(buffer)
            docx_agent = services.agents.create(
                AgentInput(
                    name="HR DOCX",
                    instructions="Gunakan poin bernomor untuk daftar. Jawab dari sumber aktif.",
                    provider="litellm",
                    model=model,
                )
            )
            docx = services.documents.upload(
                docx_agent.id, "cuplikan_people_policy.docx", buffer.getvalue()
            )
            assert docx.status == "ready"
            conversation = services.chat.new_conversation(docx_agent.id)
            # DOCX page=None is a real section citation, not a fabricated Word page.
            ask(
                {
                    "id": "dynamic01",
                    "agent": "hr",
                    "question": (
                        "Sebutkan tiga hal yang dibahas tim Human Capital "
                        "dalam pertemuan dengan para manajer."
                    ),
                    "file": "cuplikan_people_policy",
                    "pages": [None],
                    "required": ["kinerja", "potensi", "pengembangan", "pengakuan", "risiko"],
                },
                docx_agent,
                conversation,
            )
            ask(
                {
                    "id": "dynamic02",
                    "agent": "hr",
                    "question": "Jelaskan poin kedua tersebut.",
                    "file": "cuplikan_people_policy",
                    "pages": [None],
                    "required": ["pengembangan", "pengakuan"],
                },
                docx_agent,
                conversation,
            )
            dei = next(p for p in documents["hr"] if "diversity" in p.name)
            assert (
                services.documents.upload(docx_agent.id, dei.name, dei.read_bytes()).status
                == "ready"
            )
            services.documents.delete(docx_agent.id, docx.id)
            ask(
                {
                    "id": "dynamic03",
                    "agent": "hr",
                    "question": "Jelaskan lagi poin kedua itu.",
                    "unknown": True,
                },
                docx_agent,
                conversation,
            )
            # Establish history, edit whitespace, then exclude the old revision from rewriting.
            conversation = services.chat.new_conversation(agents["hr"].id)
            ask(
                {
                    "id": "dynamic04",
                    "agent": "hr",
                    "question": (
                        "Seberapa sering siklus Manajemen Karyawan yang Terintegrasi dilaksanakan?"
                    ),
                    "file": "people_policy",
                    "pages": [3],
                    "required": ["tahun"],
                },
                conversation=conversation,
            )
            doc = next(
                d
                for d in services.documents.list_documents(agents["hr"].id)
                if "people_policy" in d.filename
            )
            pages = services.documents.extraction(agents["hr"].id, doc.id)
            page = next(i for i, p in enumerate(pages) if p["page"] == 3)
            assert (
                services.documents.edit_page(
                    agents["hr"].id,
                    doc.id,
                    page,
                    " ".join(pages[page]["text"].split()),
                    doc.revision,
                ).status
                == "ready"
            )
            ask(
                {
                    "id": "dynamic05",
                    "agent": "hr",
                    "question": "Seberapa sering siklus tersebut dilaksanakan?",
                    "unknown": True,
                },
                conversation=conversation,
            )
            ask(
                {
                    "id": "dynamic06",
                    "agent": "hr",
                    "question": (
                        "Seberapa sering siklus Manajemen Karyawan yang Terintegrasi dilaksanakan?"
                    ),
                    "file": "people_policy",
                    "pages": [3],
                    "required": ["tahun"],
                }
            )

            # Fixed verifier inputs make the candidates independent of generated answer quality.
            verifier = CitationVerifier(ProviderRegistry(run_settings), run_settings)
            annual = Citation(
                tag="S1",
                chunk_id="fixed",
                document_id="fixed",
                filename="people_policy.pdf",
                page=3,
                text=text,
                score=1,
            )
            anti_path = next(p for p in documents["legal"] if "anti_corruption" in p.name)
            gifts = Citation(
                tag="S1",
                chunk_id="fixed",
                document_id="fixed",
                filename=anti_path.name,
                page=2,
                text=PdfReader(anti_path).pages[1].extract_text(),
                score=1,
            )
            fixed = [
                (
                    "v01",
                    "Siklus Manajemen Karyawan yang Terintegrasi dilaksanakan setiap tahun. [S1]",
                    annual,
                    True,
                ),
                (
                    "v02",
                    "Siklus Manajemen Karyawan yang Terintegrasi dilaksanakan setiap bulan. [S1]",
                    annual,
                    False,
                ),
                ("v03", "Karyawan memperoleh cuti tahunan 60 hari. [S1]", annual, False),
                (
                    "v04",
                    "Hadiah diperbolehkan jika sesuai batasan, prosedur, "
                    "dan hukum yang berlaku. [S1]",
                    gifts,
                    True,
                ),
                (
                    "v05",
                    "Semua hadiah kepada pihak ketiga dilarang tanpa pengecualian. [S1]",
                    gifts,
                    False,
                ),
                (
                    "v06",
                    "Perjalanan pejabat pemerintah memerlukan persetujuan "
                    "Legal and Compliance terlebih dahulu. [S1]",
                    gifts,
                    True,
                ),
            ]
            for case_id, answer, citation, expected in fixed:
                start = time.perf_counter()
                accepted = verifier.verify("Apa ketentuan kebijakan ini?", answer, [citation])
                report["verifier_records"].append(
                    {
                        "model": model,
                        "repeat": repeat + 1,
                        "id": case_id,
                        "accepted": accepted,
                        "expected": expected,
                        "pass": accepted == expected,
                        "seconds": time.perf_counter() - start,
                    }
                )
                save()

    report["summary"] = {}
    for model in MODELS:
        records = [r for r in report["records"] if r["model"] == model]
        verifier = [r for r in report["verifier_records"] if r["model"] == model]
        positive = [r for r in records if not r["expected_unknown"]]
        negative = [r for r in records if r["expected_unknown"]]
        latencies = sorted(r["seconds"] for r in records if r["provider_calls"])
        report["summary"][model] = {
            "rubric_pass": sum(r["checks"]["pass"] for r in records),
            "cases": len(records),
            "answerable_pass": sum(r["checks"]["pass"] for r in positive),
            "answerable_cases": len(positive),
            "unknown_pass": sum(r["checks"]["pass"] for r in negative),
            "unknown_cases": len(negative),
            "unverified": sum(r.get("answer") == UNVERIFIED for r in records),
            "provider_errors": sum("error_type" in r for r in records),
            "median_seconds_with_provider_call": statistics.median(latencies),
            "max_seconds_with_provider_call": max(latencies),
            "fixed_verifier_pass": sum(r["pass"] for r in verifier),
            "fixed_verifier_cases": len(verifier),
        }
    save()
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2), flush=True)
    print(f"DONE: {root / 'results.json'}", flush=True)


if __name__ == "__main__":
    logging.getLogger("httpx").setLevel(logging.WARNING)
    main()
