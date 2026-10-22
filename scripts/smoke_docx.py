"""Opt-in DOCX and numbered-follow-up smoke test using local Ollama only.

The in-memory DOCX fixture contains a public PTVI People Policy excerpt from the
prepared PDF. It is a format conversion for testing, not an official native DOCX.
No fictional company policy is added. Run after scripts.prepare_demo.
"""

import io
import json
from uuid import uuid4

from docx import Document
from pypdf import PdfReader

from src.bootstrap import initialize
from src.config import Settings
from src.knowledge.embeddings import LocalEmbeddings
from src.models import AgentInput
from src.services.chat import UNKNOWN


def main():
    settings = Settings()
    if not settings.litellm_model.startswith(("ollama/", "ollama_chat/")):
        raise ValueError("This smoke test only permits local Ollama models.")
    source = settings.data_dir / "demo/hr/people_policy_ptvi_pol_0014_id.pdf"
    text = PdfReader(source).pages[2].extract_text()
    excerpt = text.split("6. Tata Kelola Karyawan", 1)[1].split("7. Ketentuan Umum", 1)[0]
    document = Document()
    document.add_heading("Tata Kelola Karyawan", 1)
    document.add_paragraph(" ".join(excerpt.split()))
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Jenis sumber"
    table.rows[0].cells[1].text = "Kebijakan publik PT Vale Indonesia"
    buffer = io.BytesIO()
    document.save(buffer)
    scratch = settings.data_dir / "smoke-docx" / uuid4().hex
    services = initialize(settings.model_copy(update={"data_dir": scratch}))
    embeddings = LocalEmbeddings(settings)
    services.documents.embeddings = services.chat.embeddings = embeddings
    agent = services.agents.create(
        AgentInput(
            name="HR PT Vale DOCX smoke",
            instructions=(
                "Untuk permintaan daftar, jawab dengan poin bernomor. Gunakan dokumen aktif."
            ),
            provider="litellm",
            model=settings.litellm_model,
        )
    )
    doc = services.documents.upload(
        agent.id, "cuplikan_tata_kelola_pt_vale_demo.docx", buffer.getvalue()
    )
    assert doc.status == "ready", doc.error_message
    chat = services.chat.new_conversation(agent.id)
    records = []
    for question in [
        "Sebutkan tiga hal yang dibahas tim Human Capital dalam pertemuan dengan para manajer.",
        "Jelaskan poin kedua tersebut.",
    ]:
        answer = services.chat.ask(agent.id, chat, question)
        records.append(
            {
                "question": question,
                "answer": answer["content"],
                "citations": [
                    {k: c[k] for k in ("filename", "page", "locator")} for c in answer["citations"]
                ],
            }
        )
        print(json.dumps(records[-1], ensure_ascii=False), flush=True)
        assert answer["citations"] and all(c["page"] is None for c in answer["citations"])
        assert all("Tata Kelola Karyawan" in c["locator"] for c in answer["citations"])
    assert "pengembangan" in records[1]["answer"].lower()
    # Another document stays ready; the rewriter must still exclude the removed answer.
    pdf = settings.data_dir / "demo/hr/diversity_inclusion_ptvi_pol_0036_id.pdf"
    assert services.documents.upload(agent.id, pdf.name, pdf.read_bytes()).status == "ready"
    services.documents.delete(agent.id, doc.id)
    deleted = services.chat.ask(agent.id, chat, "Jelaskan lagi poin kedua itu.")
    assert deleted["content"] == UNKNOWN and not deleted["citations"]
    assert len(services.documents.ready_ids(agent.id)) == 1
    records.append({"question": "follow-up after DOCX deletion", "answer": deleted["content"]})
    (scratch / "results.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    print(f"PASS: DOCX, numbered follow-up, and deletion. Report: {scratch / 'results.json'}")


if __name__ == "__main__":
    main()
