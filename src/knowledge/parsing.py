import io
import re
import zipfile
from dataclasses import dataclass

from pypdf import PdfReader, PdfWriter

from src.knowledge.ocr import OCRConfig, OCRError, recognize_pages


class DocumentError(ValueError):
    """Safe, user-facing ingestion failure."""


@dataclass(frozen=True)
class PageText:
    page: int | None
    text: str
    locator: str | None = None


def normalize_pdf_text(text: str) -> str:
    """Join visual line wraps while retaining paragraphs and common list boundaries."""
    text = text.replace("\x00", " ").replace("\r", "")
    text = re.sub(r"(?<=[A-Za-zÀ-ÿ])-\n(?=[A-Za-zÀ-ÿ])", "", text)
    text = re.sub(
        r"(?<![.!?:;])\n(?!\n|\s*[•●▪◦]|\s*\(?\d+[.)]|\s*\([ivx]+\))",
        " ",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_pdf(data: bytes, *, ocr: OCRConfig | None = None) -> list[PageText]:
    if not data:
        raise DocumentError("File kosong. Unggah PDF yang memiliki teks.")
    if not data.lstrip().startswith(b"%PDF-"):
        raise DocumentError("File bukan PDF yang valid atau file rusak.")
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentError("PDF dilindungi kata sandi. Unggah salinan yang tidak terenkripsi.")
        if len(reader.pages) > 300:
            raise DocumentError("PDF melebihi batas 300 halaman. Pecah menjadi beberapa dokumen.")
        pages = []
        scan_numbers = [
            number
            for number, page in enumerate(reader.pages, 1)
            if sum(c.isalnum() for c in (page.extract_text() or "")) < 20 and len(page.images)
        ]
        recognized = recognize_pages(data, scan_numbers, ocr) if ocr and scan_numbers else {}
        total = 0
        for number, page in enumerate(reader.pages, start=1):
            text = recognized.get(number, page.extract_text() or "")
            # PDF extractors often put visual line wraps into the text stream.
            text = normalize_pdf_text(text)
            total += len(text)
            if total > 1_000_000:
                raise DocumentError("Teks PDF terlalu besar. Pecah menjadi beberapa dokumen.")
            if text:
                locator = f"Halaman PDF {number} · OCR" if number in recognized else None
                pages.append(PageText(number, text, locator))
        if not pages or sum(c.isalnum() for p in pages for c in p.text) < 20:
            raise DocumentError(
                "PDF tidak memiliki teks yang dapat dibaca (kosong atau hasil scan). "
                "Aktifkan OCR dan periksa setup bahasa, "
                "atau unggah PDF dengan teks yang dapat diseleksi."
            )
        return pages
    except OCRError as error:
        raise DocumentError(str(error)) from None
    except DocumentError:
        raise
    except Exception:
        raise DocumentError(
            "PDF rusak atau teksnya tidak dapat diproses. Coba file lain."
        ) from None


def pdf_page_preview(data: bytes, page_number: int) -> bytes:
    """Render a single-page PDF for the viewer without modifying the stored original."""
    try:
        reader = PdfReader(io.BytesIO(data))
        if not 1 <= page_number <= len(reader.pages):
            raise ValueError("Invalid page")
        writer = PdfWriter()
        writer.add_page(reader.pages[page_number - 1])
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()
    except Exception:
        raise DocumentError("Halaman PDF tidak dapat ditampilkan. Unduh dokumen asli.") from None


def chunk_pages(pages: list[PageText], tokenizer, size: int, overlap: int) -> list[PageText]:
    """Use model token offsets, preserving the original text and page locator."""
    if not 0 <= overlap < size:
        raise ValueError("Invalid chunk overlap")
    chunks = []
    for page in pages:
        offsets = tokenizer(
            page.text,
            add_special_tokens=False,
            return_offsets_mapping=True,
            truncation=False,
            verbose=False,
        )["offset_mapping"]
        for start in range(0, len(offsets), size - overlap):
            window = offsets[start : start + size]
            text = page.text[window[0][0] : window[-1][1]].strip()
            if text:
                chunks.append(PageText(page.page, text, page.locator))
            if start + size >= len(offsets):
                break
    return chunks


def parse_docx(data: bytes) -> list[PageText]:
    """Extract body paragraphs and table rows in order; DOCX has no stable page numbers."""
    if not data:
        raise DocumentError("File DOCX kosong. Unggah dokumen yang memiliki teks.")
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 5000
                or sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024
            ):
                raise DocumentError("Isi DOCX terlalu besar. Pecah menjadi beberapa dokumen.")
            if "word/document.xml" not in archive.namelist():
                raise DocumentError("File bukan DOCX yang valid.")
        document = Document(io.BytesIO(data))
        units, buffer, headings = [], [], []
        paragraph_number = table_number = total = 0
        first_paragraph = 0

        def section():
            return " > ".join(headings) or "Dokumen"

        def flush():
            if buffer:
                locator = f"Bagian: {section()} · paragraf {first_paragraph}–{paragraph_number}"
                units.append(PageText(None, section() + "\n" + "\n".join(buffer), locator))
                buffer.clear()

        def cell_text(cell, depth=0):
            text = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
            if depth < 5:
                for nested in cell.tables:
                    for row in nested.rows:
                        text.append(" | ".join(cell_text(c, depth + 1) for c in row.cells))
            return " ".join(text)

        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip().replace("\x00", " ")
                style = block.style.name if block.style else ""
                match = re.match(r"Heading (\d+)$", style)
                if text and (match or style == "Title"):
                    flush()
                    level = int(match[1]) if match else 1
                    headings[:] = headings[: level - 1] + [text]
                paragraph_number += 1
                if text:
                    if not buffer:
                        first_paragraph = paragraph_number
                    buffer.append(text)
                    total += len(text)
            elif isinstance(block, Table):
                flush()
                table_number += 1
                rows = [" | ".join(cell_text(cell) for cell in row.cells) for row in block.rows]
                populated = [
                    (i, row) for i, row in enumerate(rows, 1) if row.replace("|", "").strip()
                ]
                # Keep the first nonempty row as context, rather than retrieving an empty header.
                # Its full text is retained even when the table has no semantic header row.
                context = populated[0] if len(populated) > 1 else None
                for number, row in populated[1:] if context else populated:
                    header = context[1] + "\n" if context else ""
                    locator = f"Bagian: {section()} · tabel {table_number}, baris {number}"
                    if context:
                        locator += f" (konteks baris {context[0]})"
                    units.append(PageText(None, section() + "\n" + header + row, locator))
                    total += len(header) + len(row)
            if total > 1_000_000:
                raise DocumentError("Teks DOCX terlalu besar. Pecah menjadi beberapa dokumen.")
        flush()
        if sum(c.isalnum() for unit in units for c in unit.text) < 20:
            raise DocumentError(
                "DOCX tidak memiliki teks yang dapat dibaca (kosong atau hanya gambar). "
                "Unggah dokumen dengan paragraf atau tabel teks."
            )
        return units
    except DocumentError:
        raise
    except Exception:
        raise DocumentError(
            "DOCX rusak, terenkripsi, atau bukan dokumen Word yang didukung. Coba file lain."
        ) from None
