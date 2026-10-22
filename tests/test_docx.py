import io
import re
import zipfile

import pytest
from docx import Document

from src.knowledge.parsing import DocumentError, chunk_pages, parse_docx


def make_docx():
    document = Document()
    document.add_heading("Kebijakan pengembangan karyawan", 0)
    document.add_heading("Program mentoring", 1)
    document.add_paragraph("Karyawan mengikuti mentoring bersama pembimbing setiap bulan.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kegiatan"
    table.rows[0].cells[1].text = "Penanggung jawab"
    row = table.add_row()
    row.cells[0].text = "Mentoring"
    row.cells[1].text = "Tim Pembelajaran"
    document.add_heading("Evaluasi", 1)
    document.add_paragraph("Hasil pengembangan karyawan dievaluasi setiap semester.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_preserves_section_table_order_and_locators():
    units = parse_docx(make_docx())
    texts = "\n".join(unit.text for unit in units)
    assert texts.index("setiap bulan") < texts.index("Tim Pembelajaran") < texts.index("semester")
    mentoring = next(unit for unit in units if "setiap bulan" in unit.text)
    assert mentoring.page is None
    assert "Program mentoring" in mentoring.locator and "paragraf" in mentoring.locator
    table_row = next(unit for unit in units if "Tim Pembelajaran" in unit.text)
    assert "tabel 1, baris 2" in table_row.locator
    assert "Penanggung jawab" in table_row.text

    def tokenizer(text, **kwargs):
        return {"offset_mapping": [(m.start(), m.end()) for m in re.finditer(r"\S+", text)]}

    chunks = chunk_pages([mentoring], tokenizer, 4, 1)
    assert len(chunks) > 1 and all(c.locator == mentoring.locator for c in chunks)


@pytest.mark.parametrize("data", [b"", b"not a docx", b"PKbroken"])
def test_invalid_docx_has_safe_error(data):
    with pytest.raises(DocumentError):
        parse_docx(data)


def test_empty_docx_and_unrelated_zip_are_rejected():
    document, stream = Document(), io.BytesIO()
    document.save(stream)
    with pytest.raises(DocumentError, match="tidak memiliki teks"):
        parse_docx(stream.getvalue())
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("unrelated.txt", "not a Word document")
    with pytest.raises(DocumentError, match="bukan DOCX"):
        parse_docx(stream.getvalue())
